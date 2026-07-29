"""Run in-container:
    docker cp tests/test_detect_island_specs.py sqlbot:/tmp/
    docker exec sqlbot sh -c "cd /opt/sqlbot/app && .venv/bin/python -m pytest /tmp/test_detect_island_specs.py -q"
"""
import pandas as pd
from openpyxl import Workbook

from apps.datasource.utils.excel import detect_island_specs


def test_single_clean_table_yields_region_none(tmp_path):
    p = tmp_path / "one.xlsx"
    pd.DataFrame({"name": ["a", "b"], "age": [1, 2]}).to_excel(str(p), index=False)
    specs = detect_island_specs(str(p))
    assert len(specs) == 1
    assert specs[0]["region"] is None
    names = [f["fieldName"] for f in specs[0]["fields"]]
    assert "name" in names and "age" in names


def test_two_stacked_tables_yield_two_region_specs(tmp_path):
    wb = Workbook()
    ws = wb.active
    for row in [["Name", "Age"], ["Alice", "30"], ["Bob", "25"],
                [], [],
                ["City", "Pop"], ["NYC", "8"], ["LA", "4"]]:
        ws.append(row)
    p = tmp_path / "two.xlsx"
    wb.save(str(p))

    specs = detect_island_specs(str(p))
    assert len(specs) == 2
    assert [s["regionIndex"] for s in specs] == [0, 1]
    assert all(s["region"] is not None for s in specs)

    names0 = [str(f["fieldName"]) for f in specs[0]["fields"]]
    names1 = [str(f["fieldName"]) for f in specs[1]["fields"]]
    assert "Name" in names0 and "Age" in names0
    assert "City" in names1 and "Pop" in names1


def test_region_header_offset_resolved_absolute(tmp_path):
    wb = Workbook()
    ws = wb.active
    for row in [["Sales", "Q1"], ["x", "1"],
                [],
                ["MONTHLY REPORT", ""],     # title row inside the 2nd band
                ["Month", "Total"], ["Jan", "5"], ["Feb", "7"]]:
        ws.append(row)
    p = tmp_path / "offset.xlsx"
    wb.save(str(p))

    specs = detect_island_specs(str(p))
    assert len(specs) == 2
    names1 = [str(f["fieldName"]) for f in specs[1]["fields"]]
    assert "Month" in names1 and "Total" in names1


def test_clean_table_with_trailing_blank_rows_stays_region_none(tmp_path):
    # A clean single table padded with trailing blank rows must NOT be pushed to
    # the region path (which would lose dtype inference) — it stays region=None.
    wb = Workbook()
    ws = wb.active
    for row in [["name", "age"], ["a", "1"], ["b", "2"], [], [], []]:
        ws.append(row)
    p = tmp_path / "trailing.xlsx"
    wb.save(str(p))

    specs = detect_island_specs(str(p))
    assert len(specs) == 1
    assert specs[0]["region"] is None


def test_header_only_sheet_returns_single_region_none_spec(tmp_path):
    # Degenerate (no data rows) sheet must not crash and yields one region=None spec.
    p = tmp_path / "headeronly.xlsx"
    pd.DataFrame(columns=["a", "b", "c"]).to_excel(str(p), index=False)

    specs = detect_island_specs(str(p))
    assert len(specs) == 1
    assert specs[0]["region"] is None
    assert [str(f["fieldName"]) for f in specs[0]["fields"]] == ["a", "b", "c"]
